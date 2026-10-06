from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


STOPWORDS = {
    "le", "la", "les", "un", "une", "des", "de", "du", "dans", "sur", "pour", "par", "avec", "et", "ou",
    "est", "sont", "ce", "cette", "ces", "nous", "vous", "ils", "elles", "je", "tu", "quel", "quelle", "quels",
    "quelles", "comment", "pourquoi", "peut", "pouvez", "faire", "avoir", "être", "plus", "moins", "au", "aux",
}


@dataclass
class Chunk:
    source: str
    text: str


class DocumentStore:
    def __init__(self):
        self.chunks: list[Chunk] = []
        self.files: list[str] = []

    def clear(self) -> None:
        self.chunks.clear()
        self.files.clear()

    def add_files(self, paths: list[str]) -> None:
        for raw in paths:
            path = Path(raw)
            try:
                text = self._extract(path)
            except Exception:
                continue
            if not text.strip():
                continue
            self.files.append(str(path))
            for para in self._split(text):
                self.chunks.append(Chunk(source=path.name, text=para))

    def retrieve(self, query: str, limit: int = 4) -> list[Chunk]:
        qtokens = self._tokens(query)
        if not qtokens:
            return []
        scored: list[tuple[float, Chunk]] = []
        for chunk in self.chunks:
            ctokens = self._tokens(chunk.text)
            overlap = qtokens & ctokens
            if not overlap:
                continue
            score = len(overlap) / max(1, len(qtokens))
            score += 0.15 * sum(chunk.text.lower().count(t) for t in overlap)
            scored.append((score, chunk))
        scored.sort(key=lambda x: x[0], reverse=True)
        return [chunk for _, chunk in scored[:limit]]

    @staticmethod
    def _tokens(text: str) -> set[str]:
        words = re.findall(r"[A-Za-zÀ-ÿ0-9][A-Za-zÀ-ÿ0-9'’-]{2,}", text.lower())
        return {w for w in words if w not in STOPWORDS}

    @staticmethod
    def _split(text: str, max_chars: int = 900) -> list[str]:
        paragraphs = [p.strip() for p in re.split(r"\n\s*\n|(?<=[.!?])\s+", text) if p.strip()]
        out: list[str] = []
        current = ""
        for p in paragraphs:
            if len(current) + len(p) + 1 <= max_chars:
                current = (current + " " + p).strip()
            else:
                if current:
                    out.append(current)
                current = p[:max_chars]
        if current:
            out.append(current)
        return out

    def _extract(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md", ".csv"}:
            return path.read_text(encoding="utf-8", errors="ignore")
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            return "\n\n".join((p.extract_text() or "") for p in reader.pages)
        if suffix == ".docx":
            doc = Document(str(path))
            return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        if suffix in {".xlsx", ".xlsm"}:
            wb = load_workbook(str(path), read_only=True, data_only=True)
            lines: list[str] = []
            for ws in wb.worksheets:
                lines.append(f"Feuille: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    vals = [str(v) for v in row if v not in (None, "")]
                    if vals:
                        lines.append(" | ".join(vals))
            return "\n".join(lines)
        raise ValueError(f"Format non pris en charge: {suffix}")
