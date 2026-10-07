from __future__ import annotations

import html
import threading
import traceback
from datetime import datetime
from pathlib import Path

from PySide6.QtCore import QObject, Qt, Signal
from PySide6.QtWidgets import (
    QApplication,
    QCheckBox,
    QComboBox,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)
from docx import Document

from .config import Settings, app_path
from .meeting import MeetingController, QAItem, TranscriptItem


class Bridge(QObject):
    transcript = Signal(object)
    question = Signal(str)
    answer = Signal(str)
    answer_delta = Signal(str)
    qa_archived = Signal(object)
    status = Signal(str)
    error = Signal(str)
    connected = Signal(str)
    summary = Signal(str)
    auth_finished = Signal()
    summary_finished = Signal()


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.settings = Settings.load()
        self.controller = MeetingController(self.settings)
        self.bridge = Bridge()
        self._wire_backend()
        self.setWindowTitle("HENO — Meeting Assistant")
        self.resize(1320, 820)
        self.setMinimumSize(1050, 680)
        self.current_answer = ""
        self.current_summary = ""
        self._build_ui()
        self._apply_style()
        self._refresh_auth_state()
        self._refresh_archive_count()

    def _wire_backend(self) -> None:
        self.controller.on_transcript = lambda item: self.bridge.transcript.emit(item)
        self.controller.on_question = lambda q: self.bridge.question.emit(q)
        self.controller.on_answer = lambda a: self.bridge.answer.emit(a)
        self.controller.on_answer_delta = lambda d: self.bridge.answer_delta.emit(d)
        self.controller.on_qa_archived = lambda qa: self.bridge.qa_archived.emit(qa)
        self.controller.on_status = lambda s: self.bridge.status.emit(s)
        self.controller.on_error = lambda e: self.bridge.error.emit(e)
        self.bridge.transcript.connect(self._append_transcript)
        self.bridge.question.connect(self._show_question)
        self.bridge.answer.connect(self._finish_answer)
        self.bridge.answer_delta.connect(self._append_answer_delta)
        self.bridge.qa_archived.connect(self._qa_archived)
        self.bridge.status.connect(self._set_status)
        self.bridge.error.connect(self._show_error)
        self.bridge.connected.connect(self._connected)
        self.bridge.summary.connect(self._summary_ready)
        self.bridge.auth_finished.connect(lambda: self.auth_button.setEnabled(True))
        self.bridge.summary_finished.connect(lambda: self.summary_button.setEnabled(True))

    def _build_ui(self) -> None:
        root = QWidget()
        root.setObjectName("root")
        outer = QHBoxLayout(root)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(0)

        sidebar = QFrame()
        sidebar.setObjectName("sidebar")
        sidebar.setFixedWidth(270)
        sl = QVBoxLayout(sidebar)
        sl.setContentsMargins(24, 28, 24, 24)
        sl.setSpacing(10)

        logo = QLabel("HENO")
        logo.setObjectName("logo")
        sub = QLabel("MEETING NOTEBOOK")
        sub.setObjectName("subtitle")
        sl.addWidget(logo)
        sl.addWidget(sub)
        sl.addSpacing(18)

        chapter = QLabel("COMPTE")
        chapter.setObjectName("chapter")
        sl.addWidget(chapter)
        self.auth_button = QPushButton("Continuer avec ChatGPT")
        self.auth_button.clicked.connect(self.connect_chatgpt)
        sl.addWidget(self.auth_button)
        self.auth_label = QLabel("Non connecté")
        self.auth_label.setWordWrap(True)
        self.auth_label.setObjectName("muted")
        sl.addWidget(self.auth_label)

        sl.addSpacing(12)
        chapter2 = QLabel("RÉUNION")
        chapter2.setObjectName("chapter")
        sl.addWidget(chapter2)

        self.meeting_title = QLineEdit()
        self.meeting_title.setPlaceholderText("Titre de la réunion")
        sl.addWidget(self.meeting_title)

        self.language_combo = QComboBox()
        self.language_combo.addItem("Auto — Français / English", "auto")
        self.language_combo.addItem("Français", "fr")
        self.language_combo.addItem("English", "en")
        idx = self.language_combo.findData(self.settings.language)
        self.language_combo.setCurrentIndex(idx if idx >= 0 else 0)
        self.language_combo.currentIndexChanged.connect(self._save_settings)
        sl.addWidget(self.language_combo)

        self.model_combo = QComboBox()
        self.model_combo.addItems(["tiny", "base", "small", "medium"])
        self.model_combo.setCurrentText(self.settings.whisper_model)
        self.model_combo.currentTextChanged.connect(self._save_settings)
        sl.addWidget(self.model_combo)

        self.mic_check = QCheckBox("Microphone")
        self.mic_check.setChecked(self.settings.use_microphone)
        self.mic_check.stateChanged.connect(self._save_settings)
        self.system_check = QCheckBox("Son de l'ordinateur")
        self.system_check.setChecked(self.settings.use_system_audio)
        self.system_check.stateChanged.connect(self._save_settings)
        sl.addWidget(self.mic_check)
        sl.addWidget(self.system_check)

        self.docs_button = QPushButton("Documents de contexte")
        self.docs_button.clicked.connect(self.add_documents)
        sl.addWidget(self.docs_button)
        self.docs_label = QLabel("0 document")
        self.docs_label.setObjectName("muted")
        sl.addWidget(self.docs_label)

        sl.addSpacing(12)
        chapter3 = QLabel("ARCHIVES")
        chapter3.setObjectName("chapter")
        sl.addWidget(chapter3)
        self.archive_button = QPushButton("Questions & réponses")
        self.archive_button.clicked.connect(self.show_archive)
        sl.addWidget(self.archive_button)
        self.archive_label = QLabel("0 question archivée")
        self.archive_label.setObjectName("muted")
        sl.addWidget(self.archive_label)
        sl.addStretch(1)

        self.summary_button = QPushButton("Générer le compte rendu")
        self.summary_button.clicked.connect(self.generate_summary)
        sl.addWidget(self.summary_button)
        self.export_button = QPushButton("Exporter Word")
        self.export_button.clicked.connect(self.export_word)
        sl.addWidget(self.export_button)

        content = QWidget()
        content.setObjectName("desk")
        cl = QVBoxLayout(content)
        cl.setContentsMargins(28, 22, 28, 26)
        cl.setSpacing(15)

        top = QHBoxLayout()
        title_wrap = QVBoxLayout()
        book_title = QLabel("Carnet de réunion")
        book_title.setObjectName("bookTitle")
        book_subtitle = QLabel("Écouter · comprendre · répondre · archiver")
        book_subtitle.setObjectName("bookSubtitle")
        title_wrap.addWidget(book_title)
        title_wrap.addWidget(book_subtitle)
        top.addLayout(title_wrap)
        top.addStretch(1)
        self.status_dot = QLabel("●")
        self.status_dot.setObjectName("statusDot")
        self.status_label = QLabel("Prêt")
        self.status_label.setObjectName("statusText")
        top.addWidget(self.status_dot)
        top.addWidget(self.status_label)
        self.start_button = QPushButton("Démarrer la réunion")
        self.start_button.setObjectName("primary")
        self.start_button.clicked.connect(self.toggle_meeting)
        top.addWidget(self.start_button)
        cl.addLayout(top)

        pages = QHBoxLayout()
        pages.setSpacing(18)

        left_page = QFrame()
        left_page.setObjectName("paper")
        ll = QVBoxLayout(left_page)
        ll.setContentsMargins(26, 24, 26, 24)
        page_no = QLabel("PAGE 01  ·  TRANSCRIPTION")
        page_no.setObjectName("pageNo")
        ll.addWidget(page_no)
        left_title = QLabel("Ce qui se dit")
        left_title.setObjectName("pageTitle")
        ll.addWidget(left_title)
        left_rule = QFrame()
        left_rule.setObjectName("rule")
        left_rule.setFixedHeight(1)
        ll.addWidget(left_rule)
        self.transcript_view = QTextEdit()
        self.transcript_view.setObjectName("manuscript")
        self.transcript_view.setReadOnly(True)
        self.transcript_view.setPlaceholderText("La transcription apparaîtra ici, comme des notes prises au fil de la réunion…")
        ll.addWidget(self.transcript_view, 1)

        right_page = QFrame()
        right_page.setObjectName("paper")
        rl = QVBoxLayout(right_page)
        rl.setContentsMargins(26, 24, 26, 24)
        page_no2 = QLabel("PAGE 02  ·  HENO")
        page_no2.setObjectName("pageNo")
        rl.addWidget(page_no2)
        right_title = QLabel("Question & réponse proposée")
        right_title.setObjectName("pageTitle")
        rl.addWidget(right_title)
        right_rule = QFrame()
        right_rule.setObjectName("rule")
        right_rule.setFixedHeight(1)
        rl.addWidget(right_rule)

        qlabel = QLabel("QUESTION DÉTECTÉE")
        qlabel.setObjectName("marginNote")
        rl.addWidget(qlabel)
        self.question_view = QTextEdit()
        self.question_view.setObjectName("noteBox")
        self.question_view.setReadOnly(True)
        self.question_view.setMaximumHeight(160)
        self.question_view.setPlaceholderText("Une question détectée apparaîtra ici.")
        rl.addWidget(self.question_view)

        alabel = QLabel("RÉPONSE PROPOSÉE")
        alabel.setObjectName("marginNote")
        rl.addWidget(alabel)
        self.answer_view = QTextEdit()
        self.answer_view.setObjectName("noteBox")
        self.answer_view.setReadOnly(True)
        self.answer_view.setPlaceholderText("HENO proposera ici une formulation directement utilisable.")
        rl.addWidget(self.answer_view, 1)

        actions = QHBoxLayout()
        self.copy_button = QPushButton("Copier")
        self.copy_button.clicked.connect(self.copy_answer)
        actions.addWidget(self.copy_button)
        actions.addStretch(1)
        self.session_archive_label = QLabel("0 Q&R dans cette réunion")
        self.session_archive_label.setObjectName("muted")
        actions.addWidget(self.session_archive_label)
        rl.addLayout(actions)

        pages.addWidget(left_page, 3)
        pages.addWidget(right_page, 2)
        cl.addLayout(pages, 1)

        outer.addWidget(sidebar)
        outer.addWidget(content, 1)
        self.setCentralWidget(root)

    def _apply_style(self) -> None:
        self.setStyleSheet("""
            #root, QMainWindow { background: #e8e3d9; }
            QWidget { color: #2e2b27; font-family: 'Segoe UI'; font-size: 14px; }
            #sidebar { background: #f1ede5; border-right: 1px solid #cfc7b9; }
            #desk { background: #ddd7cc; }
            #logo { font-family: 'Georgia'; font-size: 34px; font-weight: 700; color: #273c34; letter-spacing: 3px; }
            #subtitle { color: #7f7669; font-size: 10px; letter-spacing: 2.5px; }
            #chapter { color: #8c7760; font-size: 10px; font-weight: 700; letter-spacing: 1.8px; margin-top: 4px; }
            #muted { color: #847c72; font-size: 12px; }
            #bookTitle { font-family: 'Georgia'; font-size: 25px; font-weight: 700; color: #292621; }
            #bookSubtitle { color: #7d7468; font-size: 12px; }
            #statusDot { color: #4f7a68; font-size: 14px; }
            #statusText { color: #665f56; margin-right: 8px; }
            #paper { background: #fffdf8; border: 1px solid #cfc7b8; border-radius: 4px; }
            #pageNo { color: #9b8d7d; font-size: 10px; letter-spacing: 1.6px; }
            #pageTitle { font-family: 'Georgia'; font-size: 23px; font-weight: 700; color: #2f2a25; margin-bottom: 5px; }
            #rule { background: #d8cfc0; border: none; margin-bottom: 8px; }
            #marginNote { color: #8a735b; font-size: 10px; font-weight: 700; letter-spacing: 1.4px; margin-top: 8px; }
            QTextEdit#manuscript { background: transparent; border: none; color: #36312c; font-family: 'Georgia'; font-size: 15px; line-height: 1.5; padding: 4px 2px; }
            QTextEdit#noteBox { background: #faf6ed; border: 1px solid #ddd3c3; border-radius: 3px; color: #34302b; font-family: 'Georgia'; font-size: 14px; padding: 10px; }
            QLineEdit, QComboBox { background: #fffdf8; border: 1px solid #c9c0b1; border-radius: 4px; padding: 8px; color: #332f2b; }
            QComboBox::drop-down { border: none; width: 24px; }
            QPushButton { background: #ebe5da; border: 1px solid #c8bfae; border-radius: 4px; padding: 9px 11px; color: #3d3832; }
            QPushButton:hover { background: #e1d9cb; }
            QPushButton#primary { background: #385f50; color: #fffdf8; border: 1px solid #385f50; font-weight: 700; padding: 11px 16px; }
            QPushButton#primary:hover { background: #2e5144; }
            QCheckBox { spacing: 8px; color: #514b44; }
        """)

    def _save_settings(self) -> None:
        if not hasattr(self, "model_combo"):
            return
        self.settings.whisper_model = self.model_combo.currentText()
        self.settings.language = self.language_combo.currentData() or "auto"
        self.settings.use_microphone = self.mic_check.isChecked()
        self.settings.use_system_audio = self.system_check.isChecked()
        self.settings.save()

    def _refresh_auth_state(self) -> None:
        if self.controller.auth.is_connected:
            self.auth_label.setText(f"Connecté : {self.settings.chatgpt_email or 'ChatGPT'}")
            self.auth_button.setText("Reconnecter ChatGPT")
        else:
            self.auth_label.setText("Non connecté — Ollama sera utilisé s'il est disponible")
            self.auth_button.setText("Continuer avec ChatGPT")

    def _refresh_archive_count(self) -> None:
        count = len(self.controller.archive.list_records(limit=5000))
        self.archive_label.setText(f"{count} question{'s' if count != 1 else ''} archivée{'s' if count != 1 else ''}")
        session_count = len(self.controller.qa_history)
        self.session_archive_label.setText(f"{session_count} Q&R dans cette réunion")

    def connect_chatgpt(self) -> None:
        self.auth_button.setEnabled(False)
        self.auth_label.setText("Ouverture de la connexion ChatGPT…")

        def task():
            try:
                result = self.controller.auth.connect()
                try:
                    self.controller.chatgpt.ensure_model()
                except Exception:
                    pass
                self.bridge.connected.emit(result.get("email") or "ChatGPT")
            except Exception as exc:
                try:
                    app_path("heno-error.log").write_text(traceback.format_exc(), encoding="utf-8")
                except Exception:
                    pass
                self.bridge.error.emit(f"{exc}\n\nJournal détaillé : %APPDATA%\\HENO-Meeting-Assistant\\heno-error.log")
            finally:
                self.bridge.auth_finished.emit()

        threading.Thread(target=task, daemon=True).start()

    def _connected(self, email: str) -> None:
        self.auth_label.setText(f"Connecté : {email}")
        self.auth_button.setText("Reconnecter ChatGPT")
        self._set_status("ChatGPT connecté")

    def toggle_meeting(self) -> None:
        self._save_settings()
        self.controller.meeting_title = self.meeting_title.text().strip()
        if not self.controller.running:
            try:
                self.transcript_view.clear()
                self.question_view.clear()
                self.answer_view.clear()
                self.current_answer = ""
                self.current_summary = ""
                self.controller.start()
                self._refresh_archive_count()
                self.start_button.setText("Arrêter la réunion")
            except Exception as exc:
                self._show_error(str(exc))
        else:
            self.controller.stop()
            self.start_button.setText("Démarrer la réunion")

    def add_documents(self) -> None:
        files, _ = QFileDialog.getOpenFileNames(
            self,
            "Ajouter des documents de contexte",
            "",
            "Documents (*.pdf *.docx *.xlsx *.xlsm *.txt *.md *.csv)",
        )
        if files:
            self.controller.documents.add_files(files)
            count = len(self.controller.documents.files)
            self.docs_label.setText(f"{count} document{'s' if count > 1 else ''}")
            self._set_status("Documents de contexte chargés")

    def _append_transcript(self, item: TranscriptItem) -> None:
        self.transcript_view.append(
            f"<span style='color:#9a8d7e'>{html.escape(item.time)}</span> "
            f"<b>{html.escape(item.source)}</b> — {html.escape(item.text)}"
        )

    def _show_question(self, question: str) -> None:
        self.question_view.setPlainText(question)
        self.answer_view.clear()
        self.current_answer = ""
        self._set_status("Question détectée — HENO prépare une réponse")

    def _append_answer_delta(self, delta: str) -> None:
        self.current_answer += delta
        self.answer_view.setPlainText(self.current_answer)
        cursor = self.answer_view.textCursor()
        cursor.movePosition(cursor.MoveOperation.End)
        self.answer_view.setTextCursor(cursor)

    def _finish_answer(self, answer: str) -> None:
        self.current_answer = answer
        self.answer_view.setPlainText(answer)
        self._set_status("Réponse prête et archivée")

    def _qa_archived(self, qa: QAItem) -> None:
        self._refresh_archive_count()

    def copy_answer(self) -> None:
        QApplication.clipboard().setText(self.answer_view.toPlainText())
        self._set_status("Réponse copiée")

    def show_archive(self) -> None:
        records = self.controller.archive.list_records(limit=500)
        dialog = QDialog(self)
        dialog.setWindowTitle("Archives HENO — Questions & réponses")
        dialog.resize(820, 620)
        layout = QVBoxLayout(dialog)
        title = QLabel("Archives des questions et réponses proposées")
        title.setObjectName("pageTitle")
        layout.addWidget(title)
        view = QTextEdit()
        view.setReadOnly(True)
        view.setObjectName("noteBox")
        if not records:
            view.setPlainText("Aucune question archivée pour le moment.")
        else:
            blocks = []
            for record in reversed(records):
                when = record.get("asked_at", "")
                try:
                    when = datetime.fromisoformat(when.replace("Z", "+00:00")).strftime("%d/%m/%Y %H:%M")
                except Exception:
                    pass
                blocks.append(
                    f"{record.get('meeting_title') or 'Réunion sans titre'}  ·  {when}\n"
                    f"QUESTION\n{record.get('question', '')}\n\n"
                    f"RÉPONSE HENO\n{record.get('answer', '')}\n"
                    + ("─" * 70)
                )
            view.setPlainText("\n\n".join(blocks))
        layout.addWidget(view, 1)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Close)
        buttons.rejected.connect(dialog.reject)
        buttons.accepted.connect(dialog.accept)
        layout.addWidget(buttons)
        dialog.exec()

    def generate_summary(self) -> None:
        if not self.controller.transcript:
            self._show_error("Aucune transcription disponible.")
            return
        self.controller.meeting_title = self.meeting_title.text().strip()
        self.summary_button.setEnabled(False)
        self._set_status("Génération du compte rendu…")

        def task():
            try:
                summary = self.controller.build_summary()
                self.bridge.summary.emit(summary)
            except Exception as exc:
                self.bridge.error.emit(str(exc))
            finally:
                self.bridge.summary_finished.emit()

        threading.Thread(target=task, daemon=True).start()

    def _summary_ready(self, summary: str) -> None:
        self.current_summary = summary
        self._set_status("Compte rendu prêt")
        dialog = QDialog(self)
        dialog.setWindowTitle("Compte rendu HENO")
        dialog.resize(760, 620)
        layout = QVBoxLayout(dialog)
        title = QLabel("Compte rendu de réunion")
        title.setObjectName("pageTitle")
        layout.addWidget(title)
        view = QTextEdit()
        view.setReadOnly(True)
        view.setObjectName("noteBox")
        view.setPlainText(summary)
        layout.addWidget(view, 1)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Close)
        buttons.rejected.connect(dialog.reject)
        buttons.accepted.connect(dialog.accept)
        layout.addWidget(buttons)
        dialog.exec()

    def export_word(self) -> None:
        if not self.controller.transcript and not self.current_summary:
            self._show_error("Rien à exporter pour le moment.")
            return
        suggested = (self.meeting_title.text().strip() or "Compte-rendu-HENO").replace("/", "-") + ".docx"
        path, _ = QFileDialog.getSaveFileName(self, "Exporter le compte rendu", suggested, "Word (*.docx)")
        if not path:
            return
        doc = Document()
        doc.add_heading(self.meeting_title.text().strip() or "Compte rendu de réunion", level=0)
        doc.add_paragraph("Généré avec HENO — Meeting Assistant")
        lang_label = {"fr": "Français", "en": "English", "auto": "Auto FR/EN"}.get(self.settings.language, "Auto FR/EN")
        doc.add_paragraph(f"Mode de transcription : {lang_label}")

        if self.current_summary:
            doc.add_heading("Compte rendu structuré", level=1)
            for line in self.current_summary.splitlines():
                if line.strip():
                    doc.add_paragraph(line.strip())

        if self.controller.qa_history:
            doc.add_heading("Questions et réponses proposées par HENO", level=1)
            for index, qa in enumerate(self.controller.qa_history, 1):
                p = doc.add_paragraph()
                p.add_run(f"Question {index} : ").bold = True
                p.add_run(qa.question)
                p2 = doc.add_paragraph()
                p2.add_run("Réponse proposée : ").bold = True
                p2.add_run(qa.answer)

        doc.add_heading("Transcription", level=1)
        for item in self.controller.transcript:
            p = doc.add_paragraph()
            p.add_run(f"[{item.time}] {item.source}: ").bold = True
            p.add_run(item.text)
        doc.save(path)
        self._set_status(f"Exporté : {Path(path).name}")

    def _set_status(self, message: str) -> None:
        self.status_label.setText(message)

    def _show_error(self, message: str) -> None:
        self._set_status("Erreur")
        QMessageBox.critical(self, "HENO", message)

    def closeEvent(self, event):
        if self.controller.running:
            self.controller.stop()
        super().closeEvent(event)
